# -*- coding: utf-8 -*-
import openpyxl
from pathlib import Path

for f in ["SMIT file.xlsx", "SMHB file.xlsx"]:
    p = Path(r"D:\TCMAI\archive\deprecated\symmap") / f"SymMap v2.0, {f}"
    print(f"\n=== {f} ===")
    if not p.exists():
        print("文件不存在")
        continue
    wb = openpyxl.load_workbook(p, read_only=True)
    ws = wb.active
    print(f"行数: {ws.max_row}, 列数: {ws.max_column}")
    for i, row in enumerate(ws.iter_rows(max_row=3)):
        print([c.value for c in row])
    wb.close()