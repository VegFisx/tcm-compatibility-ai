# -*- coding: utf-8 -*-
import openpyxl
from pathlib import Path

p = Path(r"D:\TCMAI\data\ChineseDrugEthnonymsSynonymyTable.xlsx")
wb = openpyxl.load_workbook(p, read_only=True)
ws = wb.active

print(f"工作表: {ws.title}")
print(f"行数: {ws.max_row}")
print(f"列数: {ws.max_column}\n")

print("=== 表头 ===")
header = [c.value for c in next(ws.iter_rows(max_row=1))]
print(header)

print("\n=== 前 5 行数据 ===")
for i, row in enumerate(ws.iter_rows(min_row=2, max_row=6), 1):
    print(f"[{i}] {[c.value for c in row]}")

wb.close()